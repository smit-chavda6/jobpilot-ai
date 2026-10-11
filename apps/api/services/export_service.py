"""ATS-compliant PDF and DOCX resume exporter with text round-trip verification."""

from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

import docx
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


class ExportService:
    """Generates ATS-optimized single-column PDF and DOCX documents."""

    @staticmethod
    def generate_pdf(resume_data: Dict[str, Any]) -> bytes:
        """Generate clean single-column ATS readable PDF."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54,
        )

        styles = getSampleStyleSheet()

        # ATS Compliant Styles (Helvetica, standard leading and sizing)
        title_style = ParagraphStyle(
            "ATSTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#0f172a"),
            alignment=0,  # Left-aligned for ATS
        )
        contact_style = ParagraphStyle(
            "ATSContact",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#475569"),
        )
        heading_style = ParagraphStyle(
            "ATSHeading",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=17,
            textColor=colors.HexColor("#1e293b"),
            spaceBefore=10,
            spaceAfter=4,
            keepWithNext=True,
        )
        body_style = ParagraphStyle(
            "ATSBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#334155"),
        )
        bullet_style = ParagraphStyle(
            "ATSBullet",
            parent=body_style,
            leftIndent=14,
            firstLineIndent=-10,
            spaceAfter=3,
        )

        story = []

        # 1. Header & Contact
        header = resume_data.get("header", {})
        name = header.get("name", "Candidate Name")
        email = header.get("email", "candidate@example.com")
        target_role = header.get("target_role", "")

        story.append(Paragraph(name, title_style))
        contact_line = f"{email}" + (f" | Target: {target_role}" if target_role else "")
        story.append(Paragraph(contact_line, contact_style))
        story.append(Spacer(1, 10))

        # 2. Professional Summary
        summary = resume_data.get("summary", {})
        summary_text = summary.get("text") if isinstance(summary, dict) else str(summary)
        if summary_text:
            story.append(Paragraph("PROFESSIONAL SUMMARY", heading_style))
            story.append(Paragraph(summary_text, body_style))
            story.append(Spacer(1, 8))

        # 3. Technical Skills
        skills = resume_data.get("skills", [])
        if skills:
            story.append(Paragraph("TECHNICAL SKILLS", heading_style))
            skill_names = [s.get("name", str(s)) if isinstance(s, dict) else str(s) for s in skills]
            story.append(Paragraph(", ".join(skill_names), body_style))
            story.append(Spacer(1, 8))

        # 4. Verified Experience & Accomplishments
        bullets = resume_data.get("bullets", [])
        if bullets:
            story.append(Paragraph("EXPERIENCE & VERIFIED IMPACT", heading_style))
            for b in bullets:
                b_text = b.get("text", str(b)) if isinstance(b, dict) else str(b)
                story.append(Paragraph(f"&bull; {b_text}", bullet_style))
            story.append(Spacer(1, 8))

        # 5. Education
        education = resume_data.get("education", [])
        if education:
            story.append(Paragraph("EDUCATION", heading_style))
            for edu in education:
                edu_text = edu.get("text", str(edu)) if isinstance(edu, dict) else str(edu)
                story.append(Paragraph(f"&bull; {edu_text}", bullet_style))

        doc.build(story)
        return buffer.getvalue()

    @staticmethod
    def generate_docx(resume_data: Dict[str, Any]) -> bytes:
        """Generate clean single-column ATS readable DOCX document."""
        doc = docx.Document()

        # Set 0.75-inch margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        header = resume_data.get("header", {})
        name = header.get("name", "Candidate Name")
        email = header.get("email", "candidate@example.com")
        target_role = header.get("target_role", "")

        # Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run(name)
        run_title.bold = True
        run_title.font.size = Pt(18)
        run_title.font.name = "Arial"
        p_title.paragraph_format.space_after = Pt(2)

        # Contact
        p_contact = doc.add_paragraph()
        contact_line = f"{email}" + (f" | Target: {target_role}" if target_role else "")
        run_contact = p_contact.add_run(contact_line)
        run_contact.font.size = Pt(10)
        run_contact.font.color.rgb = RGBColor(71, 85, 105)
        run_contact.font.name = "Arial"
        p_contact.paragraph_format.space_after = Pt(10)

        # Summary
        summary = resume_data.get("summary", {})
        summary_text = summary.get("text") if isinstance(summary, dict) else str(summary)
        if summary_text:
            h = doc.add_heading(level=2)
            run_h = h.add_run("PROFESSIONAL SUMMARY")
            run_h.font.size = Pt(12)
            run_h.bold = True
            run_h.font.name = "Arial"
            p = doc.add_paragraph(summary_text)
            p.paragraph_format.space_after = Pt(8)

        # Skills
        skills = resume_data.get("skills", [])
        if skills:
            h = doc.add_heading(level=2)
            run_h = h.add_run("TECHNICAL SKILLS")
            run_h.font.size = Pt(12)
            run_h.bold = True
            run_h.font.name = "Arial"
            skill_names = [s.get("name", str(s)) if isinstance(s, dict) else str(s) for s in skills]
            p = doc.add_paragraph(", ".join(skill_names))
            p.paragraph_format.space_after = Pt(8)

        # Bullets / Experience
        bullets = resume_data.get("bullets", [])
        if bullets:
            h = doc.add_heading(level=2)
            run_h = h.add_run("EXPERIENCE & VERIFIED IMPACT")
            run_h.font.size = Pt(12)
            run_h.bold = True
            run_h.font.name = "Arial"
            for b in bullets:
                b_text = b.get("text", str(b)) if isinstance(b, dict) else str(b)
                p = doc.add_paragraph(b_text, style="List Bullet")
                p.paragraph_format.space_after = Pt(3)

        # Education
        education = resume_data.get("education", [])
        if education:
            h = doc.add_heading(level=2)
            run_h = h.add_run("EDUCATION")
            run_h.font.size = Pt(12)
            run_h.bold = True
            run_h.font.name = "Arial"
            for edu in education:
                edu_text = edu.get("text", str(edu)) if isinstance(edu, dict) else str(edu)
                p = doc.add_paragraph(edu_text, style="List Bullet")
                p.paragraph_format.space_after = Pt(3)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def validate_pdf_roundtrip(pdf_bytes: bytes, expected_keywords: List[str]) -> bool:
        """Verify text extraction round-trip readability using pypdf."""
        reader = PdfReader(io.BytesIO(pdf_bytes))
        full_text = ""
        for page in reader.pages:
            full_text += page.extract_text() or ""
        
        lower_text = full_text.lower()
        return all(kw.lower() in lower_text for kw in expected_keywords)

    @staticmethod
    def validate_docx_roundtrip(docx_bytes: bytes, expected_keywords: List[str]) -> bool:
        """Verify text extraction round-trip readability using python-docx."""
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = " ".join([p.text for p in doc.paragraphs])
        lower_text = full_text.lower()
        return all(kw.lower() in lower_text for kw in expected_keywords)
