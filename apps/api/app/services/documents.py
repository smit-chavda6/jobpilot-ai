"""Resume Studio: tailored resumes and cover letters generated ONLY from confirmed facts, with every
item citing `supporting_fact_ids`, validated by the grounding validator, editable, diffable and
exportable (PDF/DOCX) only when fully grounded."""

from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from typing import Any, Literal

from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

from app.ai import service as ai
from app.ai.prompts import untrusted
from app.errors import Conflict, UnprocessableError
from app.models import GeneratedDocument, Job, ProfileFact
from app.repositories import Repo
from app.security.crypto import decrypt_field
from app.services import grounding, profile_service
from app.services.skills import family, normalize_skill

SECTION_ORDER = ["summary", "experience", "project", "education", "certification", "achievement", "skill"]
SECTION_TITLES = {
    "summary": "Summary",
    "experience": "Experience",
    "project": "Projects",
    "education": "Education",
    "certification": "Certifications",
    "achievement": "Achievements",
    "skill": "Skills",
}
MAX_BULLETS = {"experience": 12, "project": 6, "achievement": 4}


class GenItem(BaseModel):
    section: str = "experience"
    kind: Literal["claim", "framing"] = "claim"
    text: str = Field(min_length=1, max_length=800)
    supporting_fact_ids: list[str] = Field(default_factory=list)


class GenItems(BaseModel):
    items: list[GenItem] = Field(max_length=60)


class LLMCheck(BaseModel):
    results: list[dict[str, Any]]


def content_hash(content: dict[str, Any]) -> str:
    canonical = json.dumps(
        {
            "header": content.get("header", {}),
            "items": [
                {k: i.get(k) for k in ("section", "kind", "text", "supporting_fact_ids")}
                for i in content.get("items", [])
            ],
        },
        sort_keys=True,
    )
    return hashlib.sha256(canonical.encode()).hexdigest()


def _fact_view(f: ProfileFact) -> grounding.FactView:
    return grounding.FactView(str(f.id), f.text, f.attributes or {}, f.status == "confirmed")


def _group_label(f: ProfileFact) -> str:
    a = f.attributes or {}
    if f.category == "experience":
        parts = [p for p in (a.get("title"), a.get("employer")) if p]
        dates = "–".join(p for p in (a.get("start"), a.get("end")) if p)
        return " — ".join(parts) + (f" ({dates})" if dates else "")
    if f.category == "project":
        return str(a.get("project") or "")
    return ""


def _job_terms(job: Job | None) -> set[str]:
    if job is None:
        return set()
    return {s.lower() for s in (job.required_skills or []) + (job.preferred_skills or [])}


def _relevance(f: ProfileFact, job_skills: set[str], title_words: set[str]) -> float:
    skills = {s.lower() for s in profile_service.fact_skills(f)}
    words = set(re.findall(r"[a-z]+", f.text.lower()))
    return 2.0 * len(skills & job_skills) + 0.2 * len(words & title_words)


def _item(section: str, text: str, fact_ids: list[str], kind: str = "claim", group: str = "") -> dict[str, Any]:
    return {
        "id": uuid.uuid4().hex[:12],
        "section": section,
        "kind": kind,
        "text": text,
        "supporting_fact_ids": fact_ids,
        "group": group,
    }


def build_resume_items(facts: list[ProfileFact], job: Job | None, tailored: bool) -> list[dict[str, Any]]:
    job_skills = _job_terms(job)
    title_words = set(re.findall(r"[a-z]+", job.title.lower())) if job else set()
    items: list[dict[str, Any]] = []
    by_cat: dict[str, list[ProfileFact]] = {}
    for f in facts:
        by_cat.setdefault(f.category, []).append(f)
    for section in SECTION_ORDER:
        group = by_cat.get(section, [])
        if not group:
            continue
        if section == "skill":
            skill_facts = (
                sorted(
                    group,
                    key=lambda f: (str(f.attributes.get("skill", f.text)).lower() not in job_skills, group.index(f)),
                )
                if tailored
                else group
            )
            names = [str(f.attributes.get("skill") or f.text.replace("Skill:", "").strip()) for f in skill_facts]
            items.append(_item("skill", ", ".join(names), [str(f.id) for f in skill_facts]))
            continue
        ordered = group
        if tailored and section in MAX_BULLETS:
            ranked = sorted(group, key=lambda f: (-_relevance(f, job_skills, title_words), group.index(f)))
            keep = {f.id for f in ranked[: MAX_BULLETS[section]]}
            if section == "experience":
                # Keep chronological employer grouping; drop only the least relevant bullets.
                ordered = [f for f in group if f.id in keep]
            else:
                ordered = [f for f in ranked if f.id in keep]
        for f in ordered:
            items.append(_item(section, f.text, [str(f.id)], group=_group_label(f)))
    return items


def _header(profile: Any, email: str) -> dict[str, str]:
    return {
        "name": profile.full_name or "",
        "headline": profile.headline or "",
        "location": profile.location or "",
        "email": email,
        "phone": decrypt_field(profile.phone_encrypted) if profile.phone_encrypted else "",
    }


async def _run_grounding(
    session: AsyncSession,
    user_id: uuid.UUID,
    items: list[dict[str, Any]],
    facts: list[ProfileFact],
    allowed_terms: set[str],
) -> dict[str, Any]:
    fact_map = {str(f.id): _fact_view(f) for f in facts}
    report = grounding.validate(
        [
            grounding.ItemInput(i["id"], i["text"], list(i.get("supporting_fact_ids") or []), i.get("kind", "claim"))
            for i in items
        ],
        fact_map,
        allowed_terms,
    )
    # Optional LLM-as-checker (only with a real provider + consent; can only add failures).
    claims = [
        {
            "item_id": i["id"],
            "text": i["text"],
            "facts": [fact_map[fid].text for fid in i.get("supporting_fact_ids", []) if fid in fact_map],
        }
        for i in items
        if i.get("kind", "claim") == "claim" and report["items"][i["id"]]["passed"]
    ]
    if claims:
        check = await ai.generate_structured(
            session,
            user_id,
            "grounding_check",
            LLMCheck,
            {"claims": untrusted("claims", json.dumps(claims))},
            fallback=lambda: LLMCheck(results=[]),
        )
        if check.value.results:
            report = grounding.merge_llm_findings(report, check.value.results)
            report["validator"] = f"deterministic@1+{check.generator}"
    return report


def keyword_report(job: Job | None, user_skills: list[str]) -> dict[str, list[str]]:
    """Classify JD keywords: already_supported (auto-included), adjacent (ask user), unsupported (never added)."""
    out: dict[str, list[str]] = {"already_supported": [], "adjacent_ask_user": [], "unsupported": []}
    if job is None:
        return out
    have = {normalize_skill(s).lower() for s in user_skills}
    fams = {family(s) for s in user_skills} - {None, "soft", "process"}
    for s in dict.fromkeys((job.required_skills or []) + (job.preferred_skills or [])):
        if normalize_skill(s).lower() in have:
            out["already_supported"].append(s)
        elif family(s) in fams:
            out["adjacent_ask_user"].append(s)
        else:
            out["unsupported"].append(s)
    return out


async def _context(
    session: AsyncSession, user_id: uuid.UUID, job_id: uuid.UUID | None
) -> tuple[list[ProfileFact], Job | None, Any]:
    job = await Repo(session, Job, user_id).get(job_id) if job_id else None  # ownership check first
    facts = await profile_service.confirmed_facts(session, user_id)
    if not facts:
        raise UnprocessableError(
            "Confirm at least one profile fact before generating documents", code="no_confirmed_facts"
        )
    profile = await profile_service.get_profile(session, user_id)
    return facts, job, profile


def _facts_json(facts: list[ProfileFact]) -> str:
    # Facts originate from resume text, so they are wrapped as data too.
    return untrusted(
        "verified_facts", json.dumps([{"id": str(f.id), "category": f.category, "text": f.text} for f in facts])
    )


def _job_block(job: Job | None) -> str:
    if job is None:
        return untrusted("job", "(general resume, no specific job)")
    return untrusted("job", f"{job.title} at {job.company}\n{job.description_text[:6000]}")


async def generate_resume(
    session: AsyncSession, user_id: uuid.UUID, email: str, job_id: uuid.UUID | None
) -> GeneratedDocument:
    facts, job, profile = await _context(session, user_id, job_id)
    base_items = build_resume_items(facts, None, tailored=False)
    det_items = build_resume_items(facts, job, tailored=job is not None)

    result = await ai.generate_structured(
        session,
        user_id,
        "tailor_resume",
        GenItems,
        {"facts": _facts_json(facts), "job": _job_block(job), "max_items": "25"},
        fallback=lambda: GenItems(
            items=[
                GenItem(section=i["section"], text=i["text"], supporting_fact_ids=i["supporting_fact_ids"])
                for i in det_items
            ]
        ),
        allowed_urls=[job.url] if job and job.url else [],
    )
    if result.generator == "deterministic":
        items = det_items
    else:
        fact_by_id = {str(f.id): f for f in facts}
        items = [
            _item(
                i.section if i.section in SECTION_TITLES else "experience",
                i.text,
                i.supporting_fact_ids,
                group=_group_label(fact_by_id[i.supporting_fact_ids[0]])
                if i.supporting_fact_ids and i.supporting_fact_ids[0] in fact_by_id
                else "",
            )
            for i in result.value.items
        ]
    header = _header(profile, email)
    content = {"header": header, "items": items}
    report = await _run_grounding(session, user_id, items, facts, set())
    doc = Repo(session, GeneratedDocument, user_id).add(
        job_id=job.id if job else None,
        kind="resume",
        title=f"Resume — {job.title} at {job.company}" if job else "Resume — general",
        content=content,
        base_content={"header": header, "items": base_items},
        grounding_report=report,
        keyword_report=keyword_report(job, await profile_service.user_skills(session, user_id)),
        grounded=report["passed"],
        generator=result.generator,
        prompt_version=result.prompt_version,
    )
    await session.flush()
    return doc


def build_cover_letter_items(facts: list[ProfileFact], job: Job | None, name: str) -> list[dict[str, Any]]:
    job_skills = _job_terms(job)
    title_words = set(re.findall(r"[a-z]+", job.title.lower())) if job else set()
    role = (
        f"the {job.title} role at {job.company}"
        if job and job.company
        else (f"the {job.title} role" if job else "this role")
    )
    items = [
        _item("letter", "Dear Hiring Team,", [], kind="framing"),
        _item("letter", f"I am writing to apply for {role}.", [], kind="framing"),
    ]
    evidence = sorted(
        (f for f in facts if f.category in ("experience", "project", "achievement")),
        key=lambda f: -_relevance(f, job_skills, title_words),
    )[:3]
    for f in evidence:
        items.append(_item("letter", f"One relevant example: {f.text.rstrip('.')}.", [str(f.id)]))
    skill_facts = (
        [f for f in facts if f.category == "skill" and str(f.attributes.get("skill", "")).lower() in job_skills]
        if job_skills
        else []
    )
    if skill_facts:
        names = [str(f.attributes.get("skill")) for f in skill_facts[:8]]
        joined = ", ".join(names[:-1]) + (f" and {names[-1]}" if len(names) > 1 else names[0])
        items.append(
            _item(
                "letter",
                f"Skills from my profile that match this posting include {joined}.",
                [str(f.id) for f in skill_facts[:8]],
            )
        )
    company = job.company if job and job.company else "your team"
    items.append(
        _item(
            "letter",
            f"Thank you for considering my application. I would welcome the chance to discuss "
            f"how I could contribute to {company}.",
            [],
            kind="framing",
        )
    )
    items.append(_item("letter", f"Sincerely, {name}".strip().rstrip(","), [], kind="framing"))
    return items


async def generate_cover_letter(
    session: AsyncSession, user_id: uuid.UUID, email: str, job_id: uuid.UUID | None
) -> GeneratedDocument:
    facts, job, profile = await _context(session, user_id, job_id)
    det_items = build_cover_letter_items(facts, job, profile.full_name or "")
    result = await ai.generate_structured(
        session,
        user_id,
        "cover_letter",
        GenItems,
        {"facts": _facts_json(facts), "job": _job_block(job)},
        fallback=lambda: GenItems(
            items=[
                GenItem(section="letter", kind=i["kind"], text=i["text"], supporting_fact_ids=i["supporting_fact_ids"])
                for i in det_items
            ]
        ),
        allowed_urls=[job.url] if job and job.url else [],
    )
    items = (
        det_items
        if result.generator == "deterministic"
        else [_item("letter", i.text, i.supporting_fact_ids, kind=i.kind) for i in result.value.items]
    )
    allowed = _allowed_framing(job, profile.full_name or "")
    report = await _run_grounding(session, user_id, items, facts, allowed)
    doc = Repo(session, GeneratedDocument, user_id).add(
        job_id=job.id if job else None,
        kind="cover_letter",
        title=f"Cover letter — {job.title} at {job.company}" if job else "Cover letter — general",
        content={"header": _header(profile, email), "items": items},
        base_content={},
        grounding_report=report,
        keyword_report={},
        grounded=report["passed"],
        generator=result.generator,
        prompt_version=result.prompt_version,
    )
    await session.flush()
    return doc


def _allowed_framing(job: Job | None, name: str) -> set[str]:
    terms = {name}
    if job:
        terms |= {job.title, job.company}
    return {t for t in terms if t}


async def update_document(
    session: AsyncSession,
    user_id: uuid.UUID,
    doc: GeneratedDocument,
    items: list[dict[str, Any]],
    header: dict[str, str] | None,
) -> GeneratedDocument:
    facts = await profile_service.confirmed_facts(session, user_id)
    clean: list[dict[str, Any]] = []
    for raw in items[:80]:
        gi = GenItem.model_validate(raw)
        clean.append(
            {
                "id": str(raw.get("id") or uuid.uuid4().hex[:12])[:24],
                "section": gi.section,
                "kind": gi.kind,
                "text": gi.text.strip(),
                "supporting_fact_ids": gi.supporting_fact_ids,
                "group": str(raw.get("group") or "")[:200],
            }
        )
    job = await Repo(session, Job, user_id).get(doc.job_id) if doc.job_id else None
    profile = await profile_service.get_profile(session, user_id)
    allowed = _allowed_framing(job, profile.full_name or "") if doc.kind == "cover_letter" else set()
    report = await _run_grounding(session, user_id, clean, facts, allowed)
    content = {"header": {**doc.content.get("header", {}), **(header or {})}, "items": clean}
    doc.content, doc.grounding_report, doc.grounded = content, report, report["passed"]
    doc.version += 1
    return doc


def diff(doc: GeneratedDocument, facts: dict[str, ProfileFact]) -> list[dict[str, Any]]:
    """Before/after view with per-line provenance (cited fact texts)."""
    base = doc.base_content.get("items", []) if doc.base_content else []
    base_by_text = {i["text"]: idx for idx, i in enumerate(base)}
    base_fact_ids = {fid for i in base for fid in i.get("supporting_fact_ids", [])}
    rows: list[dict[str, Any]] = []
    seen_texts: set[str] = set()
    for pos, item in enumerate(doc.content.get("items", [])):
        text = item["text"]
        seen_texts.add(text)
        if text in base_by_text:
            status = "unchanged" if base_by_text[text] == pos else "moved"
        elif set(item.get("supporting_fact_ids", [])) & base_fact_ids:
            status = "modified"
        else:
            status = "added"
        rows.append(
            {
                "status": status,
                "section": item["section"],
                "after": text,
                "before": base[base_by_text[text]]["text"] if text in base_by_text else None,
                "provenance": [
                    {"fact_id": fid, "text": facts[fid].text if fid in facts else "(missing fact)"}
                    for fid in item.get("supporting_fact_ids", [])
                ],
            }
        )
    for b in base:
        if b["text"] not in seen_texts:
            rows.append(
                {
                    "status": "removed",
                    "section": b["section"],
                    "after": None,
                    "before": b["text"],
                    "provenance": [
                        {"fact_id": fid, "text": facts[fid].text if fid in facts else ""}
                        for fid in b.get("supporting_fact_ids", [])
                    ],
                }
            )
    return rows


# ------------------------------------------------------------------ export


def _require_grounded(doc: GeneratedDocument) -> None:
    if not doc.grounded:
        failing = [
            {"item_id": k, "reasons": v["reasons"]}
            for k, v in doc.grounding_report.get("items", {}).items()
            if not v["passed"]
        ]
        raise Conflict(
            "Export is blocked until every item passes grounding checks",
            code="grounding_failed",
            details={"failing_items": failing},
        )


def _sections(doc: GeneratedDocument) -> list[tuple[str, list[dict[str, Any]]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in doc.content.get("items", []):
        grouped.setdefault(item["section"], []).append(item)
    order = SECTION_ORDER + [s for s in grouped if s not in SECTION_ORDER]
    return [(s, grouped[s]) for s in order if s in grouped]


def export_pdf(doc: GeneratedDocument) -> bytes:
    _require_grounded(doc)
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=doc.title,
        author=doc.content["header"].get("name", ""),
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=13)
    h = ParagraphStyle("h", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, spaceBefore=8)
    sub = ParagraphStyle("sub", parent=body, fontName="Helvetica-Bold")
    header = doc.content.get("header", {})
    flow: list[Any] = []
    if header.get("name"):
        flow.append(Paragraph(escape(header["name"]), styles["Title"]))
    contact = " | ".join(escape(header[k]) for k in ("headline", "location", "email", "phone") if header.get(k))
    if contact:
        flow.append(Paragraph(contact, body))
    if doc.kind == "cover_letter":
        flow.append(Spacer(1, 10))
        for item in doc.content.get("items", []):
            flow.append(Paragraph(escape(item["text"]), body))
            flow.append(Spacer(1, 6))
    else:
        for section, items in _sections(doc):
            flow.append(Paragraph(SECTION_TITLES.get(section, section.title()), h))
            last_group = None
            for item in items:
                if item.get("group") and item["group"] != last_group:
                    flow.append(Paragraph(escape(item["group"]), sub))
                    last_group = item["group"]
                bullet = "" if section in ("summary", "skill") else "• "
                flow.append(Paragraph(bullet + escape(item["text"]), body))
    pdf.build(flow)
    return buf.getvalue()


def export_docx(doc: GeneratedDocument) -> bytes:
    _require_grounded(doc)
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    d.core_properties.title = doc.title
    header = doc.content.get("header", {})
    if header.get("name"):
        d.add_heading(header["name"], level=0)
    contact = " | ".join(header[k] for k in ("headline", "location", "email", "phone") if header.get(k))
    if contact:
        d.add_paragraph(contact)
    if doc.kind == "cover_letter":
        for item in doc.content.get("items", []):
            d.add_paragraph(item["text"])
    else:
        for section, items in _sections(doc):
            d.add_heading(SECTION_TITLES.get(section, section.title()), level=1)
            last_group = None
            for item in items:
                if item.get("group") and item["group"] != last_group:
                    p = d.add_paragraph()
                    p.add_run(item["group"]).bold = True
                    last_group = item["group"]
                if section in ("summary", "skill"):
                    d.add_paragraph(item["text"])
                else:
                    d.add_paragraph(item["text"], style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def doc_dict(doc: GeneratedDocument) -> dict[str, Any]:
    return {
        "id": str(doc.id),
        "job_id": str(doc.job_id) if doc.job_id else None,
        "kind": doc.kind,
        "title": doc.title,
        "content": doc.content,
        "grounding_report": doc.grounding_report,
        "keyword_report": doc.keyword_report,
        "grounded": doc.grounded,
        "generator": doc.generator,
        "prompt_version": doc.prompt_version,
        "version": doc.version,
        "content_hash": content_hash(doc.content),
        "created_at": doc.created_at.isoformat(),
        "updated_at": doc.updated_at.isoformat(),
    }
