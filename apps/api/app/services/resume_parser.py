"""Resume text extraction and heuristic decomposition into *candidate* ProfileFacts.

Parsed facts are never usable downstream until the user confirms them (spec §6).
Extraction runs in a worker thread with a hard timeout so a pathological file cannot hang
the API (see ADR-0003 for why this is a thread, not an OS sandbox, on this deployment).
"""

from __future__ import annotations

import asyncio
import io
import re
from dataclasses import dataclass, field
from datetime import date
from typing import Any

from app.errors import UnprocessableError
from app.services.skills import detect_education, extract_skills, normalize_skill

PARSE_TIMEOUT_SECONDS = 20
MAX_TEXT_CHARS = 60_000

SECTION_ALIASES: dict[str, set[str]] = {
    "experience": {
        "experience",
        "work experience",
        "professional experience",
        "employment",
        "employment history",
        "work history",
        "relevant experience",
        "internships",
        "internship",
    },
    "education": {"education", "academics", "academic background", "education and training"},
    "project": {"projects", "personal projects", "academic projects", "selected projects", "key projects"},
    "skill": {
        "skills",
        "technical skills",
        "core skills",
        "technologies",
        "tools",
        "skills and tools",
        "tech stack",
        "core competencies",
    },
    "certification": {"certifications", "certificates", "licenses", "licenses and certifications"},
    "summary": {"summary", "profile", "professional summary", "objective", "about me", "about"},
    "achievement": {"achievements", "awards", "honors", "honours", "awards and achievements"},
}
_HEADING_INDEX = {alias: section for section, aliases in SECTION_ALIASES.items() for alias in aliases}

_MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
DATE_RANGE_RE = re.compile(
    rf"(?P<start>(?:{_MONTH}\s+)?(?:19|20)\d{{2}}|\d{{1,2}}/(?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*"
    rf"(?P<end>(?:{_MONTH}\s+)?(?:19|20)\d{{2}}|\d{{1,2}}/(?:19|20)\d{{2}}|present|current|now|today)",
    re.IGNORECASE,
)
BULLET_RE = re.compile(r"^\s*(?:[-•*▪‣◦●■–]|\d+[.)])\s+")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{8,}\d)")
TITLE_WORDS = re.compile(
    r"\b(engineer|developer|analyst|scientist|manager|intern|lead|designer|consultant|specialist|"
    r"architect|researcher|associate|director|officer|administrator|programmer|technician|head|"
    r"assistant|fellow|trainee|founder|executive)\b",
    re.IGNORECASE,
)


@dataclass
class ParsedFact:
    category: str
    text: str
    attributes: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedResume:
    facts: list[ParsedFact]
    contact: dict[str, str]
    estimated_years_experience: float | None
    detected_education_level: str | None


# ------------------------------------------------------------------ text extraction


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise UnprocessableError("Encrypted PDFs are not supported", code="parse_failed")
    pages = [(p.extract_text() or "") for p in reader.pages[:20]]
    return "\n".join(pages)


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(lines)


def _extract_sync(ext: str, data: bytes) -> str:
    if ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    else:
        text = data.decode("utf-8")
    return text.replace("\r\n", "\n").replace("\x00", "")[:MAX_TEXT_CHARS]


async def extract_text(ext: str, data: bytes) -> str:
    try:
        return await asyncio.wait_for(asyncio.to_thread(_extract_sync, ext, data), PARSE_TIMEOUT_SECONDS)
    except TimeoutError as exc:
        raise UnprocessableError("Parsing took too long; try a simpler file", code="parse_failed") from exc
    except UnprocessableError:
        raise
    except Exception as exc:  # malformed documents
        raise UnprocessableError("Could not read this document", code="parse_failed") from exc


# ------------------------------------------------------------------ decomposition


def _heading(line: str) -> str | None:
    cleaned = re.sub(r"[:#*_=]+", "", line).strip().lower()
    if not cleaned or len(cleaned) > 40:
        return None
    return _HEADING_INDEX.get(cleaned)


def _year(token: str) -> int | None:
    m = re.search(r"(19|20)\d{2}", token)
    return int(m.group(0)) if m else None


def _split_role_header(line: str) -> dict[str, str]:
    attrs: dict[str, str] = {}
    m = DATE_RANGE_RE.search(line)
    rest = line
    if m:
        attrs["start"], attrs["end"] = m.group("start").strip(), m.group("end").strip()
        rest = (line[: m.start()] + line[m.end() :]).strip()
    rest = re.sub(r"[()\[\]]", " ", rest).strip(" ,|-–—")
    if " at " in rest:
        title, company = rest.split(" at ", 1)
        parts = [title, company]
    else:
        parts = [p.strip() for p in re.split(r"\s+[|–—-]\s+|\s*[|–—]\s*|,\s+", rest) if p.strip()]
    if parts:
        if len(parts) >= 2 and not TITLE_WORDS.search(parts[0]) and TITLE_WORDS.search(parts[1]):
            parts[0], parts[1] = parts[1], parts[0]
        attrs["title"] = parts[0].strip()
        if len(parts) > 1:
            attrs["employer"] = parts[1].strip()
        if len(parts) > 2:
            attrs["location"] = parts[2].strip()
    return attrs


def _estimate_years(ranges: list[tuple[int, int]]) -> float | None:
    if not ranges:
        return None
    merged: list[list[int]] = []
    for s, e in sorted(ranges):
        if merged and s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    return float(sum(max(0, e - s) for s, e in merged)) or None


def decompose(text: str) -> ParsedResume:
    lines = [ln.rstrip() for ln in text.split("\n")]
    contact: dict[str, str] = {}
    if m := EMAIL_RE.search(text):
        contact["email"] = m.group(0)
    if m := PHONE_RE.search(text):
        contact["phone"] = m.group(0).strip()
    for ln in lines:
        if ln.strip():
            if not EMAIL_RE.search(ln) and len(ln.split()) <= 5 and not _heading(ln):
                contact["name"] = ln.strip()
            break

    facts: list[ParsedFact] = []
    section: str | None = None
    header: dict[str, str] = {}
    ranges: list[tuple[int, int]] = []
    summary_buf: list[str] = []
    this_year = date.today().year

    def add(category: str, txt: str, attrs: dict[str, Any] | None = None) -> None:
        txt = re.sub(r"\s+", " ", txt).strip(" -•*")
        if len(txt) >= 3:
            facts.append(ParsedFact(category, txt[:600], dict(attrs or {})))

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        sec = _heading(line)
        if sec:
            section, header = sec, {}
            continue
        if section is None:
            continue
        is_bullet = bool(BULLET_RE.match(raw))
        content = BULLET_RE.sub("", raw).strip()

        if section == "summary":
            summary_buf.append(content)
        elif section == "skill":
            body = content.split(":", 1)[1] if ":" in content and len(content.split(":", 1)[0]) < 30 else content
            for token in re.split(r"[,;|•/]|\s{2,}", body):
                token = token.strip(" .")
                if 1 <= len(token) <= 40:
                    skill = normalize_skill(token)
                    add("skill", f"Skill: {skill}", {"skill": skill})
        elif section in ("experience", "project"):
            looks_header = not is_bullet and (
                DATE_RANGE_RE.search(content)
                or (
                    len(content) < 90
                    and not content.endswith(".")
                    and (TITLE_WORDS.search(content) or section == "project")
                )
            )
            if looks_header:
                header = _split_role_header(content) if section == "experience" else {"project": content}
                if section == "project" and (dm := DATE_RANGE_RE.search(content)):
                    header = {
                        "project": DATE_RANGE_RE.sub("", content).strip(" ,|-–—()"),
                        "start": dm.group("start"),
                        "end": dm.group("end"),
                    }
                if section == "experience" and header.get("start"):
                    s = _year(header["start"])
                    end_raw = header.get("end", "")
                    e = this_year if re.match(r"present|current|now|today", end_raw, re.I) else _year(end_raw)
                    if s and e and e >= s:
                        ranges.append((s, e))
                continue
            if not is_bullet and facts and facts[-1].category == section and content[:1].islower():
                facts[-1].text = f"{facts[-1].text} {content}"[:600]
                continue
            attrs = {k: v for k, v in header.items() if v}
            add(section, content, attrs)
        elif section == "education":
            edu: dict[str, Any] = {}
            if lvl := detect_education(content):
                edu["degree_level"] = lvl
            if dm := DATE_RANGE_RE.search(content):
                edu["start"], edu["end"] = dm.group("start"), dm.group("end")
            if facts and facts[-1].category == "education" and not edu and not is_bullet:
                facts[-1].text = f"{facts[-1].text}, {content}"[:600]
                continue
            add("education", content, edu)
        else:
            add(section, content)

    if summary_buf:
        facts.insert(0, ParsedFact("summary", " ".join(summary_buf)[:600], {}))

    # Skills mentioned inside experience/project bullets are recorded as attributes (tools used).
    for f in facts:
        if f.category in ("experience", "project", "achievement"):
            tools = extract_skills(f.text)
            if tools:
                f.attributes["tools"] = tools

    edu_levels = [f.attributes.get("degree_level") for f in facts if f.category == "education"]
    order = ["phd", "master", "bachelor", "associate", "high_school"]
    detected = next((lvl for lvl in order if lvl in edu_levels), None)
    return ParsedResume(facts, contact, _estimate_years(ranges), detected)
